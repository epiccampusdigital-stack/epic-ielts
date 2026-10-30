# Budget and Spending tabs for the calendar workbook (run inside build_xlsx.py, which defines wb, header, fonts, box)
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

BLUE = Font(name=F, size=10, color='0000FF')
BLUE_B = Font(name=F, size=10, color='0000FF', bold=True)
GREEN = Font(name=F, size=10, color='008000')
YELLOW = PatternFill('solid', fgColor='FFFF00')
TINT = PatternFill('solid', fgColor='E8EEF7')
LKR = '#,##0;(#,##0);-'
PCT = '0%;(0%);-'

bu = wb.create_sheet('Budget', 1)
bu['A1'] = 'Budget: Korea Scholarship Batch 1'; bu['A1'].font = title
bu['A2'] = 'All amounts in LKR. Blue numbers are inputs you can change; black cells are formulas. Yellow cells still need a number. Record every payment on the Spending tab.'; bu['A2'].font = sub
for col, w in zip('ABCDEFGH', [34, 16, 18, 18, 18, 16, 18, 40]):
    bu.column_dimensions[col].width = w

def label(r, text, f=base):
    c = bu.cell(row=r, column=1, value=text); c.font = f; c.border = box; return c

def val(r, col, v, fmt=LKR, font=base, fill=None):
    c = bu.cell(row=r, column=col, value=v); c.font = font; c.number_format = fmt; c.border = box
    if fill: c.fill = fill
    return c

# 1. Income
bu['A4'] = '1. Income'; bu['A4'].font = bold
header(bu, 5, ['Item', 'Amount', 'Note'])
rows = [
    ('Students', 11, '0', BLUE, 'From Ishara, 28 Oct 2026: one more student added (was 10)'),
    ('Course fee per student', 125000, LKR, BLUE, 'From Ishara, 26 Oct 2026'),
    ('Total course fees', '=B6*B7', LKR, base, ''),
    ('Campus share per student', 25000, LKR, BLUE, 'From Ishara: LKR 25,000 per student goes to the campus'),
    ('Campus share total', '=B6*B9', LKR, base, ''),
    ('Left to budget', '=B8-B10', LKR, bold, 'The percentages below are applied to this amount'),
]
for i, (lab, v, fmt, fnt, note) in enumerate(rows, start=6):
    label(i, lab, bold if lab == 'Left to budget' else base)
    val(i, 2, v, fmt, fnt)
    c = bu.cell(row=i, column=3, value=note); c.font = sub; c.border = box
bu.merge_cells('C6:H6'); bu.merge_cells('C7:H7'); bu.merge_cells('C9:H9'); bu.merge_cells('C11:H11')

# 2. Allocation
bu['A13'] = '2. Where the money goes'; bu['A13'].font = bold
header(bu, 14, ['Line', '% of budget', 'Planned (LKR)', 'Your figure (LKR)', 'Spent so far (LKR)', 'Left (LKR)', 'Status'])
LINES = [('Campus', 0.10, 100000), ('Staff', 0.40, 400000), ('System', 0.12, 120000), ('Web', 0.10, 100000),
         ('Miscellaneous (transport and other)', 0.10, 100000), ('Accommodation and food', 0.30, 300000)]
first_line = 15
for i, (lab, pct, amt) in enumerate(LINES, start=first_line):
    label(i, lab)
    val(i, 2, pct, PCT, BLUE)
    val(i, 3, f'=B{i}*$B$11')
    val(i, 4, amt, LKR, BLUE)
    val(i, 5, f"=SUMIFS(Spending!$D:$D,Spending!$B:$B,A{i})", LKR, GREEN)
    val(i, 6, f'=D{i}-E{i}')
    c = bu.cell(row=i, column=7, value=f'=IF(E{i}>D{i},"Overspent",IF(E{i}>=0.9*D{i},"Nearly used","OK"))'); c.font = base; c.border = box
last_line = first_line + len(LINES) - 1
tr = last_line + 1
label(tr, 'Total', bold)
for col, fmt in [(2, PCT), (3, LKR), (4, LKR), (5, LKR), (6, LKR)]:
    L = get_column_letter(col)
    val(tr, col, f'=SUM({L}{first_line}:{L}{last_line})', fmt, bold)
bu.cell(row=first_line, column=8, value='Percentages and amounts as given by Ishara, 26 Oct 2026').font = sub

ck = tr + 1
label(ck, 'Check against the money left to budget', bold)
c = bu.cell(row=ck, column=2, value=f'=IF(ROUND(B{tr},4)=1,"Adds up to 100%","Adds up to "&TEXT(B{tr},"0%")&", not 100%")'); c.font = bold; c.border = box
bu.merge_cells(start_row=ck, start_column=2, end_row=ck, end_column=4)
c = bu.cell(row=ck + 1, column=1, value='Money left after your figures (negative = over budget)'); c.font = bold; c.border = box
val(ck + 1, 2, f'=B11-D{tr}', LKR, bold)
bu.cell(row=ck + 1, column=3, value='Your six lines total more than the money left after the campus share. Lower the % or amounts (blue) until this is zero or positive.').font = sub
red = Font(name=F, size=10, bold=True, color='D30A28')
bu.conditional_formatting.add(f'B{ck + 1}', FormulaRule(formula=[f'B{ck + 1}<0'], font=red))
bu.conditional_formatting.add(f'G{first_line}:G{last_line}', FormulaRule(formula=[f'G{first_line}="Overspent"'], font=red))
bu.conditional_formatting.add(f'B{ck}', FormulaRule(formula=[f'ROUND(B{tr},4)<>1'], font=red))

# 3. Staff costs
s0 = ck + 4
bu.cell(row=s0, column=1, value='3. Staff costs (inside the Staff line)').font = bold
header(bu, s0 + 1, ['Teacher', 'Pay basis', 'Rate (LKR)', 'Units', 'Cost (LKR)', 'Status', 'Hours in the plan', 'Note', 'Sessions (half days)'])
bu.column_dimensions['I'].width = 14
STAFF = [
    ('English Teacher 1', 'Per session', 5400, 'Agreed', 'Anna. Wed, Thu & Sat 9:00 to 12:00. LKR 1,800 an hour x 3 hours = 5,400 a session. 12 sessions = 64,800'),
    ('Transport: English Teacher 1', 'Per session', 2600, 'Agreed', 'Anna transport, LKR 2,600 a day x 12 days = 31,200'),
    ('English Teacher 2', 'Per session', 6000, 'Suggested', 'Week 1 grammar + all mock interviews. Mix of 3 h and 3.5 h sessions. Suggested LKR 6,000 a session. Change when agreed'),
    ('English Teacher 3', 'Per month', 55000, 'Confirmed', 'From Ishara: LKR 55,000 per month, confirmed. English afternoons 1:00 to 4:30 (3.5 h)'),
    ('English Teacher 4 (TBC)', 'Per session', 5500, 'Cover only', 'Only paid if they cover a session. Enter sessions covered in column I'),
    ('Korean Teacher 1', 'Staff (no extra cost)', 0, 'Staff', 'Campus staff'),
    ('Korean Teacher 2', 'Staff (no extra cost)', 0, 'Staff', 'Campus staff'),
]
import collections as _c
_b = _c.defaultdict(set)
for _s in d['sessions']:
    if _s.get('teacher'): _b[_s['teacher']].add((_s['dayNumber'], _s['startTime'] < '12:00'))
HALF = {T[k]['name']: len(v) for k, v in _b.items()}
HALF['Transport: English Teacher 1'] = HALF.get('English Teacher 1', 0)
dv = DataValidation(type='list', formula1='"Per hour,Per month,Per session,Staff (no extra cost)"', allow_blank=True)
bu.add_data_validation(dv)
months_row = s0 + 2 + len(STAFF) + 3
f0 = s0 + 2
for i, (name, basis, rate, status, note) in enumerate(STAFF, start=f0):
    label(i, name)
    c = bu.cell(row=i, column=2, value=basis); c.font = BLUE; c.border = box; dv.add(c)
    c = val(i, 3, rate, LKR, BLUE, YELLOW if rate is None else None)
    if rate is None:
        c.comment = Comment('Enter the agreed rate for this teacher.', 'EPIC')
    val(i, 7, f'=SUMIFS(Calendar!$H:$H,Calendar!$L:$L,A{i})', '0.0', GREEN)
    val(i, 4, f'=IF(B{i}="Per hour",G{i},IF(B{i}="Per month",$B${months_row},IF(B{i}="Per session",I{i},0)))', '0.0')
    val(i, 5, f'=IFERROR(C{i}*D{i},0)')
    c = bu.cell(row=i, column=6, value=status); c.font = base; c.border = box
    c = bu.cell(row=i, column=8, value=note); c.font = sub; c.border = box
    val(i, 9, HALF.get(name, 0), '0', BLUE if name.startswith('English Teacher 4') else GREEN)
fl = f0 + len(STAFF) - 1
label(fl + 1, 'Total staff cost', bold); val(fl + 1, 5, f'=SUM(E{f0}:E{fl})', LKR, bold)
label(fl + 2, 'Left in the Staff line', bold); val(fl + 2, 5, f'=D{first_line + 1}-E{fl + 1}', LKR, bold)
bu.conditional_formatting.add(f'E{fl + 2}', FormulaRule(formula=[f'E{fl + 2}<0'], font=red))
label(months_row, 'Programme length in months (for monthly pay)')
val(months_row, 2, 1, '0.0', BLUE)
bu.cell(row=months_row, column=3, value='30 teaching days = 1 month').font = sub

# 4. Handy numbers
h0 = months_row + 2
bu.cell(row=h0, column=1, value='4. Per student').font = bold
header(bu, h0 + 1, ['Item', 'Amount (LKR)'])
af = first_line + 5
items = [
    ('Food and accommodation per student per day', f'=IFERROR(D{af}/B6/30,0)'),
    ('Food and accommodation per student, whole programme', f'=IFERROR(D{af}/B6,0)'),
    ('Teaching cost per student (planned staff cost)', f'=IFERROR(E{fl + 1}/B6,0)'),
    ('Total budget per student (your figures + campus share)', f'=IFERROR((D{tr}+B10)/B6,0)'),
]
for i, (lab, f) in enumerate(items, start=h0 + 2):
    label(i, lab); val(i, 2, f)
bu.freeze_panes = 'A4'

# Spending log
sp = wb.create_sheet('Spending', 2)
sp['A1'] = 'Spending log'; sp['A1'].font = title
sp['A2'] = 'One row per payment. Choose the category from the list so the Budget tab adds it up. The first row is an example: replace or delete it.'; sp['A2'].font = sub
header(sp, 4, ['Date', 'Category', 'What it was for', 'Amount (LKR)', 'Paid by', 'Receipt / note'], [14, 34, 44, 16, 18, 30])
cats = ','.join(l[0] for l in LINES)
dv2 = DataValidation(type='list', formula1=f'"{cats}"', allow_blank=True)
sp.add_data_validation(dv2)
ex = [dt.date(2026, 10, 30), 'Accommodation and food', 'Example: rice, vegetables and gas for week 1', 18500, 'Campus office', 'Example row']
for c, v in enumerate(ex, 1):
    cell = sp.cell(row=5, column=c, value=v); cell.font = Font(name=F, size=10, italic=True, color='5B6780'); cell.border = box
sp.cell(row=5, column=1).number_format = 'DD MMM YYYY'; sp.cell(row=5, column=4).number_format = LKR
for r in range(5, 205):
    dv2.add(sp.cell(row=r, column=2))
    sp.cell(row=r, column=1).number_format = 'DD MMM YYYY'
    sp.cell(row=r, column=4).number_format = LKR
    for c in range(1, 7):
        sp.cell(row=r, column=c).border = box
        if r > 5: sp.cell(row=r, column=c).font = BLUE
sp['H4'] = 'Total spent'; sp['H4'].font = bold
sp['I4'] = '=SUM(D5:D204)'; sp['I4'].number_format = LKR; sp['I4'].font = bold
sp.column_dimensions['H'].width = 14; sp.column_dimensions['I'].width = 16
sp.freeze_panes = 'A5'
