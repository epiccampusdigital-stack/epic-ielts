import json, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

d = json.load(open('plan.json'))
T = {t['key']: t for t in d['teachers']}
F = 'Arial'
NAVY = '002761'
hdr_font = Font(name=F, bold=True, color='FFFFFF', size=10)
hdr_fill = PatternFill('solid', fgColor=NAVY)
gold_line = Side(style='medium', color='0056D2')
base = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
title = Font(name=F, size=16, bold=True, color=NAVY)
sub = Font(name=F, size=10, color='555555')
thin = Side(style='thin', color='D9D9D9')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
FILLS = {
 'ENGLISH': 'E3EEFF', 'KOREAN': 'FCE7F3', 'IT_LAB': 'ECFDF5', 'ACTIVITY': 'FFFBEB',
 'ASSESSMENT': 'F3E8FF', 'BREAK': 'F1F5F9', 'ADMIN': 'F1F5F9',
}
LABEL = {'ENGLISH':'English class','KOREAN':'Korean class','IT_LAB':'IT lab','ACTIVITY':'Activity','ASSESSMENT':'Test / interview','BREAK':'Break','ADMIN':'Admin'}
LANG = {'EN':'English','KR':'Korean',None:''}
def tname(k): return T[k]['name'] if k else ''
def tm(s):
    h,m = map(int,s.split(':')); return dt.time(h,m)
def date(s): return dt.date.fromisoformat(s)
day_date = {x['dayNumber']: x['date'] for x in d['days']}

def header(ws, row, cols, widths=None):
    for i,c in enumerate(cols,1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = center; cell.border = Border(left=thin, right=thin, top=thin, bottom=gold_line)
    if widths:
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width = w

wb = Workbook()

# ---------- Overview ----------
ov = wb.active; ov.title = 'Overview'
ov['A1'] = 'Korea Scholarship Program: 30-Day Calendar'; ov['A1'].font = title
ov['A2'] = 'Day 1 Fri 30 Oct 2026 to Day 30 Sat 28 Nov 2026 · two sessions a day: morning (8:30 to 11:30, or 9:00 to 12:00 on Anna days) and afternoon 1:00 to 4:30 · 11 students · Epic Campus'; ov['A2'].font = sub
ov['A3'] = 'Plan of 29 Oct: English Teacher 1 (Anna) teaches English 9:00 to 12:00 on Wed, Thu and Sat (12 days from 4 Nov), with Korean those afternoons. Other days: Korean mornings, English afternoons with Teacher 3. Week 1 (Days 1, 4, 5) is the grammar foundation with Teacher 2. Korean follows the 24 EStudykorea lessons and the 4 Korean workbooks. Budget and spending are on the Budget tab.'; ov['A3'].font = sub
r = 5
ov.cell(row=r, column=1, value='Teaching hours by language').font = bold
r += 1; header(ov, r, ['Language','Hours','Share'])
langs = [('English','EN'),('Korean','KR')]
for i,(lab,code) in enumerate(langs):
    rr = r+1+i
    ov.cell(row=rr, column=1, value=lab).font = base
    ov.cell(row=rr, column=2, value=f"=SUMIFS(Calendar!$H:$H,Calendar!$K:$K,\"{lab}\")").font = base
    ov.cell(row=rr, column=3, value=f"=IFERROR(B{rr}/SUM($B${r+1}:$B${r+2}),0)").font = base
    ov.cell(row=rr, column=2).number_format = '0.0'; ov.cell(row=rr, column=3).number_format = '0.0%'
    for c in range(1,4): ov.cell(row=rr, column=c).border = box
r += 4
ov.cell(row=r, column=1, value='Hours by teacher').font = bold
r += 1; header(ov, r, ['Teacher','Focus','Confirmed','Sessions','Total hours','Avg hours / day'])
first = r+1
for i,t in enumerate(d['teachers']):
    rr = first+i
    vals = [t['name'], t['focus'], 'Yes' if t['confirmed'] else 'To be confirmed']
    for c,v in enumerate(vals,1): ov.cell(row=rr, column=c, value=v).font = base
    ov.cell(row=rr, column=4, value=f'=COUNTIFS(Calendar!$L:$L,A{rr})').font = base
    ov.cell(row=rr, column=5, value=f'=SUMIFS(Calendar!$H:$H,Calendar!$L:$L,A{rr})').font = base
    ov.cell(row=rr, column=6, value=f'=E{rr}/30').font = base
    ov.cell(row=rr, column=5).number_format = '0.0'; ov.cell(row=rr, column=6).number_format = '0.0'
    for c in range(1,7): ov.cell(row=rr, column=c).border = box
r = first + len(d['teachers']) + 2
ov.cell(row=r, column=1, value='Colour key').font = bold
for i,(k,lab) in enumerate(LABEL.items()):
    c = ov.cell(row=r+1+i, column=1, value=lab); c.font = base; c.fill = PatternFill('solid', fgColor=FILLS[k]); c.border = box
r = r + len(LABEL) + 2
ov.cell(row=r, column=1, value='How to use').font = bold
notes = ['Calendar: every session, one row each. Filter by date, teacher or subject.',
         'Grid: the whole program at a glance, one column per day. Good for printing.',
         'Days: the theme and focus for each day, plus the IT lab task.',
         'One tab per teacher with only their sessions.',
         'To change a teacher, edit the Teacher column on Calendar; totals here update. The live version of this plan is the tracker on epicielts.live.']
for i,n in enumerate(notes): ov.cell(row=r+1+i, column=1, value='• '+n).font = base
for col,w in zip('ABCDEF',[30,28,16,10,12,15]): ov.column_dimensions[col].width = w

# ---------- Calendar ----------
cal = wb.create_sheet('Calendar')
cols = ['Day','Date','Weekday','Week','Start','End','Time','Hours','Session','Type','Language','Teacher','Details']
header(cal, 1, cols, [6,12,11,6,8,8,14,7,46,16,10,26,60])
sess = sorted(d['sessions'], key=lambda s:(s['dayNumber'], s['startTime']))
for i,s in enumerate(sess, start=2):
    dd = date(day_date[s['dayNumber']])
    cal.cell(row=i, column=1, value=s['dayNumber'])
    cal.cell(row=i, column=2, value=dd).number_format = 'DD MMM YYYY'
    cal.cell(row=i, column=3, value=f'=TEXT(B{i},"dddd")')
    cal.cell(row=i, column=4, value=f'=INT((A{i}-1)/6)+1')
    cal.cell(row=i, column=5, value=tm(s['startTime'])).number_format = 'h:mm AM/PM'
    cal.cell(row=i, column=6, value=tm(s['endTime'])).number_format = 'h:mm AM/PM'
    cal.cell(row=i, column=7, value=f'=TEXT(E{i},"h:mm")&" to "&TEXT(F{i},"h:mm")')
    cal.cell(row=i, column=8, value=f'=ROUND((F{i}-E{i})*24,2)').number_format = '0.00'
    cal.cell(row=i, column=9, value=s['title'])
    cal.cell(row=i, column=10, value=LABEL[s['subject']])
    cal.cell(row=i, column=11, value=LANG[s.get('language')])
    cal.cell(row=i, column=12, value=tname(s.get('teacher')))
    cal.cell(row=i, column=13, value=s.get('description') or '')
    fill = PatternFill('solid', fgColor=FILLS[s['subject']])
    for c in range(1,14):
        cell = cal.cell(row=i, column=c); cell.font = base; cell.border = box
        cell.alignment = Alignment(vertical='top', wrap_text=(c in (9,13)))
    cal.cell(row=i, column=10).fill = fill; cal.cell(row=i, column=9).fill = fill
last = len(sess)+1
tab = Table(displayName='Calendar', ref=f'A1:M{last}')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
cal.add_table(tab)
cal.freeze_panes = 'C2'
# header style again (table style may override visually)
for c in range(1,14):
    cal.cell(row=1,column=c).font = hdr_font; cal.cell(row=1,column=c).fill = hdr_fill

# ---------- Grid ----------
g = wb.create_sheet('Grid')
slots = sorted({s['startTime'] for s in sess})
g['A1'] = 'Program at a glance'; g['A1'].font = title
g.cell(row=3, column=1, value='Time').font = hdr_font; g.cell(row=3,column=1).fill = hdr_fill; g.cell(row=3,column=1).alignment = center
g.column_dimensions['A'].width = 13
for di in range(1,31):
    c = g.cell(row=3, column=di+1, value=f"Day {di}\n{date(day_date[di]).strftime('%a %d %b')}")
    c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = box
    g.column_dimensions[get_column_letter(di+1)].width = 18
g.row_dimensions[3].height = 30
by = {}
for s in sess:
    k = (s['dayNumber'], s['startTime'])
    if k in by and s['subject'] != 'BREAK':
        prev = by[k]
        by[k] = dict(prev, title=prev['title'] + ' + ' + s['title'] + ' (' + tname(s.get('teacher')) + ')')
    else:
        by[k] = s
for ri,st in enumerate(slots, start=4):
    lab = tm(st).strftime('%I:%M %p').lstrip('0')
    c = g.cell(row=ri, column=1, value=lab); c.font = bold; c.alignment = center; c.border = box
    g.row_dimensions[ri].height = 18 if all(by.get((di,st),{}).get('subject') in (None,'BREAK') for di in range(1,31)) else 58
    for di in range(1,31):
        s = by.get((di, st))
        cell = g.cell(row=ri, column=di+1); cell.border = box
        if not s: continue
        txt = s['title'] if s['subject']=='BREAK' else s['title'] + f"\n{tm(s['startTime']).strftime('%I:%M').lstrip('0')} to {tm(s['endTime']).strftime('%I:%M').lstrip('0')}" + (f"\n{tname(s.get('teacher'))}" if s.get('teacher') else '')
        cell.value = txt; cell.font = Font(name=F, size=8); cell.alignment = Alignment(wrap_text=True, vertical='top')
        cell.fill = PatternFill('solid', fgColor=FILLS[s['subject']])
g.freeze_panes = 'B4'
g.page_setup.orientation = 'landscape'; g.page_setup.fitToWidth = 0
g.sheet_properties.pageSetUpPr.fitToPage = True; g.page_setup.fitToHeight = 1

# ---------- Days ----------
ds = wb.create_sheet('Days')
header(ds, 1, ['Day','Date','Week','Theme','English focus','Korean focus (suggested)','IT lab task'], [6,14,6,26,52,30,60])
for i,x in enumerate(d['days'], start=2):
    vals = [x['dayNumber'], date(x['date']), None, x['theme'], x['englishFocus'], x['koreanFocus'], x['itLabTask']]
    for c,v in enumerate(vals,1):
        cell = ds.cell(row=i, column=c, value=v); cell.font = base; cell.alignment = wrap; cell.border = box
    ds.cell(row=i, column=2).number_format = 'ddd DD MMM'
    ds.cell(row=i, column=3, value=f'=INT((A{i}-1)/6)+1')
    ds.row_dimensions[i].height = 62
ds.freeze_panes = 'D2'

# ---------- Teacher tabs: one row per session (half day) ----------
for t in d['teachers']:
    name = t['name'].replace('(','').replace(')','')[:31]
    ws = wb.create_sheet(name)
    ws['A1'] = t['name']; ws['A1'].font = title
    ws['A2'] = f"{t['focus']} · {'Confirmed' if t['confirmed'] else 'To be confirmed'}"; ws['A2'].font = sub
    ws['A3'] = 'Sessions:'; ws['A3'].font = bold
    ws['D3'] = 'Hours:'; ws['D3'].font = bold
    mine = [s for s in sess if s.get('teacher') == t['key']]
    groups = {}
    for s in mine:
        groups.setdefault((s['dayNumber'], s['startTime'] < '12:00'), []).append(s)
    start = 5
    header(ws, start, ['Day','Date','Session','Start','End','Hours','What happens in this session','Type'], [6,14,12,10,10,7,90,16])
    for i, (key, ss) in enumerate(sorted(groups.items(), key=lambda kv: (kv[0][0], not kv[0][1])), start=start+1):
        day, am = key
        ss.sort(key=lambda x: x['startTime'])
        st, en = min(x['startTime'] for x in ss), max(x['endTime'] for x in ss)
        what = '\n'.join(f"{tm(x['startTime']).strftime('%I:%M').lstrip('0')} to {tm(x['endTime']).strftime('%I:%M').lstrip('0')}  {x['title']}" for x in ss)
        kinds = sorted({LABEL[x['subject']] for x in ss})
        main = ss[0]['subject'] if len({x['subject'] for x in ss}) == 1 else ('ASSESSMENT' if any(x['subject']=='ASSESSMENT' for x in ss) else ss[0]['subject'])
        vals = [day, date(day_date[day]), 'Morning' if am else 'Afternoon', tm(st), tm(en), None, what, ' + '.join(kinds)]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v); cell.font = base; cell.border = box
            cell.alignment = Alignment(vertical='top', wrap_text=c == 7)
        ws.cell(row=i, column=2).number_format = 'ddd DD MMM'
        ws.cell(row=i, column=4).number_format = 'h:mm AM/PM'; ws.cell(row=i, column=5).number_format = 'h:mm AM/PM'
        ws.cell(row=i, column=6, value=f'=ROUND((E{i}-D{i})*24,2)').number_format = '0.00'
        ws.cell(row=i, column=8).fill = PatternFill('solid', fgColor=FILLS[main])
        ws.row_dimensions[i].height = 15 * len(ss) + 3
    endr = start + len(groups)
    ws['B3'] = f'=COUNT(A{start+1}:A{max(endr, start+1)})'; ws['B3'].font = bold
    ws['E3'] = f'=SUM(F{start+1}:F{max(endr, start+1)})'; ws['E3'].number_format = '0.0'; ws['E3'].font = bold
    ws['G3'] = 'One row per session. Hours = start to finish, including the tea break.' + (' Mock interviews run in a side room during Teacher 3\'s afternoon.' if t['key'] == 'EN_SPEAKING' else ''); ws['G3'].font = sub
    ws.freeze_panes = f'A{start+1}'

exec(open('budget_sheet.py').read())
wb.save('/home/claude/deliverables/Korea_Program_30Day_Calendar.xlsx')
print('saved', len(sess))
